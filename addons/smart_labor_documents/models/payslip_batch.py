import base64
import html
import io
import re
import unicodedata
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from odoo import _, fields, models
from odoo.exceptions import UserError


class PayslipDocumentBatch(models.Model):
    _name = "employee.payslip.document.batch"
    _description = "Envio Masivo de Boletas"
    _inherit = ["mail.thread", "mail.activity.mixin"]
    _order = "create_date desc"

    CLIENT_HEADER_ALIASES = {
        "codigo": "employee_code",
        "apellidosynombres": "employee_name",
        "correo": "employee_email",
        "celular": "employee_phone",
        "registro_patronal": "employer_registry",
        "periodo": "period",
        "fecha_ingreso": "hire_date",
        "regimen_pensionario": "pension_regime",
        "cuspp": "cuspp",
        "cargo": "job_title",
        "entidad_bancaria": "bank_name",
        "numero_de_cuenta": "bank_account",
        "inicio_vacaciones": "vacation_start",
        "fin_vacaciones": "vacation_end",
        "dias_trabajados": "worked_days",
        "horas_efectivas_trab": "worked_hours",
        "dias_subsidiados": "subsidized_days",
        "dias_descontados": "discounted_days",
        "total_ingresos": "total_income",
        "total_descuentos": "total_discounts",
        "total_aportes": "total_employer_contributions",
        "total_a_pagar": "net_amount",
        "personal_confianza": "trusted_employee",
        "personal_direccion": "management_employee",
    }

    name = fields.Char(string="Nombre", required=True, default=lambda self: _("Envio Masivo de Boletas"))
    document_type_id = fields.Many2one(
        "employee.document.type",
        string="Documento",
        default=lambda self: self.env.ref("employee_legajo_documents.document_type_payslip", raise_if_not_found=False),
        ondelete="restrict",
    )
    template_id = fields.Many2one(
        "employee.payslip.document.template",
        string="Plantilla",
        required=True,
        ondelete="restrict",
    )
    period = fields.Char(string="Periodo", required=True, help="Ejemplo: 2026-06")
    responsible_id = fields.Many2one("res.users", string="Responsable", default=lambda self: self.env.user, required=True, readonly=True)
    description = fields.Text(string="Descripcion Envio")
    sent_date = fields.Datetime(string="Fecha Envio", readonly=True)
    excel_layout = fields.Selection(
        selection=[
            ("simple", "Simple"),
            ("classified_headers", "Cliente I/D/AE"),
        ],
        default="classified_headers",
        required=True,
        string="Formato Excel",
    )
    classification_header_row = fields.Integer(default=1, string="Fila Clasificacion")
    concept_header_row = fields.Integer(default=2, string="Fila Conceptos")
    data_start_row = fields.Integer(default=3, string="Primera Fila Datos")
    excel_file = fields.Binary(string="Archivo Excel", required=True, attachment=True)
    excel_filename = fields.Char(string="Nombre del archivo Excel")
    document_ids = fields.One2many("employee.payslip.document", "batch_id", string="Documentos generados")
    document_count = fields.Integer(string="Generados", compute="_compute_document_stats")
    sent_count = fields.Integer(string="Enviados", compute="_compute_document_stats")
    read_count = fields.Integer(string="Leidos", compute="_compute_document_stats")
    unread_count = fields.Integer(string="No leidos", compute="_compute_document_stats")
    signed_count = fields.Integer(string="Recepciones confirmadas", compute="_compute_document_stats")
    unsigned_count = fields.Integer(string="Pendientes de recepcion", compute="_compute_document_stats")
    error_count = fields.Integer(string="Errores", compute="_compute_document_stats")
    error_log = fields.Text(string="Registro de errores", readonly=True)
    state = fields.Selection(
        selection=[
            ("draft", "Borrador"),
            ("processing", "En Proceso"),
            ("processed", "Procesado"),
            ("validated", "Validado"),
            ("rejected", "Rechazado"),
            ("active", "Activo"),
            ("done", "Finalizado"),
            ("error", "Error"),
        ],
        default="draft",
        required=True,
        tracking=True,
        string="Estado",
    )

    def _compute_document_stats(self):
        for batch in self:
            documents = batch.document_ids
            batch.document_count = len(documents)
            if batch.state in ("validated", "active", "done"):
                batch.sent_count = len(documents)
                batch.read_count = len(documents.filtered("is_read"))
                batch.unread_count = batch.sent_count - batch.read_count
            else:
                batch.sent_count = 0
                batch.read_count = 0
                batch.unread_count = 0
            batch.signed_count = len(documents.filtered(lambda document: document.signature_state == "signed"))
            batch.unsigned_count = len(
                documents.filtered(lambda document: document.signature_required and document.signature_state != "signed")
            )
            batch.error_count = len([line for line in (batch.error_log or "").splitlines() if line.strip()])

    def action_process(self):
        for batch in self:
            batch.write({"state": "processing"})
            batch._process_excel()

    def action_set_done(self):
        self.write({"state": "done"})

    def action_set_active(self):
        self.write({"state": "active"})

    def action_validate_send(self):
        for batch in self:
            documents = batch.document_ids
            documents.write({"state": "published"})
            documents.action_render_template()
            documents.action_sync_legajo_document()
            batch._send_payslip_notifications()
            batch.write(
                {
                    "state": "validated",
                    "sent_date": fields.Datetime.now(),
                }
            )

    def action_reject(self):
        for batch in self:
            batch.document_ids.write({"state": "draft"})
            batch.write({"state": "rejected"})

    def action_view_documents(self):
        self.ensure_one()
        return {
            "type": "ir.actions.act_window",
            "name": _("Documentos Generados"),
            "res_model": "employee.payslip.document",
            "view_mode": "list,form",
            "domain": [("batch_id", "=", self.id)],
            "context": {"default_batch_id": self.id},
        }

    def action_view_read_documents(self):
        action = self.action_view_documents()
        action["domain"] = [("batch_id", "=", self.id), ("is_read", "=", True)]
        return action

    def action_view_unread_documents(self):
        action = self.action_view_documents()
        action["domain"] = [("batch_id", "=", self.id), ("is_read", "=", False)]
        return action

    def action_view_signed_documents(self):
        action = self.action_view_documents()
        action["domain"] = [("batch_id", "=", self.id), ("signature_state", "=", "signed")]
        return action

    def action_view_unsigned_documents(self):
        action = self.action_view_documents()
        action["domain"] = [
            ("batch_id", "=", self.id),
            ("signature_required", "=", True),
            ("signature_state", "!=", "signed"),
        ]
        return action

    def _process_excel(self):
        self.ensure_one()
        if not self.excel_file:
            raise UserError(_("Debe subir un archivo Excel."))

        try:
            file_content = base64.b64decode(self.excel_file)
            workbook = load_workbook(io.BytesIO(file_content), data_only=True)
            sheet = workbook.active
        except Exception as error:
            raise UserError(_("No se pudo leer el archivo Excel: %s") % error) from error

        rows = list(sheet.iter_rows(values_only=True))
        if self.excel_layout == "classified_headers":
            self._process_classified_headers_excel(rows)
            return

        if len(rows) < 2:
            raise UserError(_("El archivo Excel debe incluir una fila de encabezados y al menos una fila de datos."))

        headers = [self._normalize_header(value) for value in rows[0]]
        lookup_column = self._normalize_header(self.template_id.employee_lookup_column)
        if lookup_column not in headers:
            raise UserError(_("No se encontro la columna '%s' en el encabezado del Excel.") % self.template_id.employee_lookup_column)

        Employee = self.env["hr.employee"].sudo()
        Document = self.env["employee.payslip.document"].sudo()
        errors = []
        created_documents = self.env["employee.payslip.document"]

        for row_number, row in enumerate(rows[1:], start=2):
            values = {
                headers[index]: self._format_cell(row[index] if index < len(row) else "")
                for index in range(len(headers))
                if headers[index]
            }
            values = self._apply_header_aliases(values)
            if not any(values.values()):
                continue

            lookup_value = values.get(lookup_column)
            if not lookup_value:
                errors.append(_("Fila %(row)s: falta el valor para identificar al empleado.") % {"row": row_number})
                continue

            employee = self._find_employee(Employee, lookup_value)
            if not employee:
                errors.append(
                    _("Fila %(row)s: no se encontro empleado usando %(field)s = %(value)s.")
                    % {
                        "row": row_number,
                        "field": self.template_id.employee_lookup_field,
                        "value": lookup_value,
                    }
                )
                continue

            values.setdefault("employee_name", employee.name)
            values.setdefault("period", self.period)

            document_name = "%s - %s" % (employee.name, self.period)
            existing = Document.search(
                [
                    ("employee_id", "=", employee.id),
                    ("period", "=", self.period),
                    ("template_id", "=", self.template_id.id),
                ],
                limit=1,
            )
            document_values = {
                "name": document_name,
                "template_id": self.template_id.id,
                "batch_id": self.id,
                "employee_id": employee.id,
                "period": self.period,
                "values_json": values,
                "digital_signer_user_id": self.template_id.digital_signer_user_id.id,
                "state": "draft",
            }
            if existing:
                existing.write(document_values)
                existing.action_render_template()
                created_documents |= existing
            else:
                document = Document.create(document_values)
                document.action_render_template()
                created_documents |= document

        self.write(
            {
                "state": "error" if errors and not created_documents else "processed",
                "error_log": "\n".join(errors),
            }
        )
        if not created_documents and errors:
            raise UserError(_("No se generaron documentos.\n\n%s") % "\n".join(errors))

    def _find_employee(self, Employee, lookup_value):
        self.ensure_one()
        field_name = self.template_id.employee_lookup_field
        operator = "=ilike" if field_name in ("work_email", "name") else "="
        return Employee.search([(field_name, operator, str(lookup_value).strip())], limit=1)

    def _process_classified_headers_excel(self, rows):
        self.ensure_one()
        if len(rows) < self.data_start_row:
            raise UserError(_("El archivo Excel no tiene suficientes filas para el formato I/D/AE configurado."))

        classification_index = self.classification_header_row - 1
        concept_index = self.concept_header_row - 1
        data_start_index = self.data_start_row - 1
        classification_row = rows[classification_index]
        concept_row = rows[concept_index]

        headers = [self._normalize_header(value) for value in concept_row]
        lookup_column = self._normalize_header(self.template_id.employee_lookup_column)
        if lookup_column not in headers:
            raise UserError(_("No se encontro la columna '%s' en la fila de conceptos.") % self.template_id.employee_lookup_column)

        Employee = self.env["hr.employee"].sudo()
        Document = self.env["employee.payslip.document"].sudo()
        errors = []
        created_documents = self.env["employee.payslip.document"]

        for row_number, row in enumerate(rows[data_start_index:], start=self.data_start_row):
            values = {
                headers[index]: self._format_cell(row[index] if index < len(row) else "")
                for index in range(len(headers))
                if headers[index]
            }
            values = self._apply_header_aliases(values)
            if not any(values.values()):
                continue

            lookup_value = values.get(lookup_column)
            if not lookup_value:
                errors.append(_("Fila %(row)s: falta el valor para identificar al empleado.") % {"row": row_number})
                continue

            employee = self._find_employee(Employee, lookup_value)
            if not employee:
                errors.append(
                    _("Fila %(row)s: no se encontro empleado usando %(field)s = %(value)s.")
                    % {
                        "row": row_number,
                        "field": self.template_id.employee_lookup_field,
                        "value": lookup_value,
                    }
                )
                continue

            values.setdefault("employee_name", employee.name)
            values.setdefault("period", self.period)
            classified_values = self._build_classified_values(classification_row, concept_row, row)
            values.update(classified_values)

            document_name = "%s - %s" % (employee.name, self.period)
            existing = Document.search(
                [
                    ("employee_id", "=", employee.id),
                    ("period", "=", self.period),
                    ("template_id", "=", self.template_id.id),
                ],
                limit=1,
            )
            document_values = {
                "name": document_name,
                "template_id": self.template_id.id,
                "batch_id": self.id,
                "employee_id": employee.id,
                "period": self.period,
                "values_json": values,
                "digital_signer_user_id": self.template_id.digital_signer_user_id.id,
                "state": "draft",
            }
            if existing:
                existing.write(document_values)
                existing.action_render_template()
                created_documents |= existing
            else:
                document = Document.create(document_values)
                document.action_render_template()
                created_documents |= document

        self.write(
            {
                "state": "error" if errors and not created_documents else "processed",
                "error_log": "\n".join(errors),
            }
        )
        if not created_documents and errors:
            raise UserError(_("No se generaron documentos.\n\n%s") % "\n".join(errors))

    def _build_classified_values(self, classification_row, concept_row, row):
        lines_by_section = {
            "I": [],
            "D": [],
            "AE": [],
        }
        for index, concept in enumerate(concept_row):
            section = str(classification_row[index] or "").strip().upper() if index < len(classification_row) else ""
            if section not in lines_by_section:
                continue
            amount = row[index] if index < len(row) else ""
            if self._is_empty_cell(amount) or self._is_zero_amount(amount):
                continue
            label = str(concept or "").strip()
            if not label:
                continue
            lines_by_section[section].append(
                {
                    "label": label,
                    "amount": self._format_amount_cell(amount),
                }
            )

        return {
            "income_lines": lines_by_section["I"],
            "discount_lines": lines_by_section["D"],
            "employer_contribution_lines": lines_by_section["AE"],
            "income_lines_html": self._lines_to_html(lines_by_section["I"]),
            "discount_lines_html": self._lines_to_html(lines_by_section["D"]),
            "employer_contribution_lines_html": self._lines_to_html(lines_by_section["AE"]),
        }

    @staticmethod
    def _is_empty_cell(value):
        return value is None or str(value).strip() == ""

    @staticmethod
    def _is_zero_amount(value):
        if value is None or isinstance(value, bool):
            return False
        if isinstance(value, (int, float, Decimal)):
            return Decimal(str(value)) == Decimal("0")
        value_text = str(value).strip()
        if not value_text:
            return False
        normalized = value_text.replace(",", "")
        try:
            return Decimal(normalized) == Decimal("0")
        except InvalidOperation:
            return False

    @staticmethod
    def _lines_to_html(lines):
        if not lines:
            return ""
        return "".join(
            '<tr><td>%s</td><td class="amount">%s</td></tr>'
            % (html.escape(str(line["label"])), html.escape(str(line["amount"])))
            for line in lines
        )

    @staticmethod
    def _format_amount_cell(value):
        if value is None:
            return ""
        if isinstance(value, bool):
            return str(value)
        if isinstance(value, (int, float, Decimal)):
            return "%.2f" % float(value)
        value_text = str(value).strip()
        if not value_text:
            return ""
        normalized = value_text.replace(",", "")
        try:
            return "%.2f" % float(Decimal(normalized))
        except (InvalidOperation, ValueError):
            return value_text

    def _send_payslip_notifications(self):
        Mail = self.env["mail.mail"].sudo()
        base_url = self.env["ir.config_parameter"].sudo().get_param("web.base.url", "").rstrip("/")
        for document in self.document_ids.filtered(lambda item: item.state == "published"):
            email_to = document.employee_id.work_email or document.employee_id.user_id.email
            if not email_to:
                continue
            url = "%s/my/payslip-documents/%s" % (base_url, document.id)
            body = """
                <p>Hola %s,</p>
                <p>Se ha publicado tu boleta correspondiente al periodo <strong>%s</strong>.</p>
                <p><a href="%s">Ver boleta</a></p>
            """ % (document.employee_id.name, document.period, url)
            Mail.create(
                {
                    "subject": "Boleta de pago %s" % document.period,
                    "body_html": body,
                    "email_to": email_to,
                    "auto_delete": False,
                }
            )

    @staticmethod
    def _normalize_header(value):
        text = str(value or "").strip().lower()
        text = unicodedata.normalize("NFKD", text)
        text = "".join(char for char in text if not unicodedata.combining(char))
        text = re.sub(r"[^a-z0-9]+", "_", text)
        return text.strip("_")

    def _apply_header_aliases(self, values):
        mapped_values = dict(values)
        for source_key, target_key in self.CLIENT_HEADER_ALIASES.items():
            if source_key in values and target_key not in mapped_values:
                mapped_values[target_key] = values[source_key]
        return mapped_values

    @staticmethod
    def _format_cell(value):
        if value is None:
            return ""
        if hasattr(value, "strftime"):
            return value.strftime("%Y-%m-%d")
        return value
